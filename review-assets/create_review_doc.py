from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / 'Invaria Tomorrow Review.docx'
FONT = Path('C:/Windows/Fonts/arial.ttf')
FONT_BOLD = Path('C:/Windows/Fonts/arialbd.ttf')

NAVY = '11121A'
INDIGO = '4F46E5'
INK = '161824'
MUTED = '60677C'
LINE = 'DFE3EE'
SOFT = 'EEEDFF'
WHITE = 'FFFFFF'

def pc(value):
    return f'#{value}'

PHASES = [
    ('01', 'Ingestion', 'Bring in the repository and any business documents.',
     'Repository link and documents', 'Validate source, collect supported files, remove build output', 'Code and document inventory',
     'For TaskForge the repository import completed. New external repositories still need broader verification.'),
    ('02', 'Understand the code', 'Turn supported source files into useful application facts.',
     'Supported JavaScript and TypeScript', 'Find routes, handlers, inputs, data lookups, and access checks', 'Code structure map',
     'TaskForge maps 19 endpoints. Parsing coverage is still limited to supported patterns and files.'),
    ('03', 'Map the connections', 'Show how a request reaches functions and data.',
     'Routes, functions, inputs, queries', 'Link the reachable code paths into a security graph', 'Interactive request to data path',
     'The demo graph makes one endpoint path easy to inspect without reading raw source first.'),
    ('04', 'Find potential risks', 'Apply static checks to identify patterns that need attention.',
     'Mapped routes and source evidence', 'Run supported authorization and code-risk checks', 'Potential findings with severity and evidence',
     'TaskForge currently shows 16 potential risks. These are review candidates, not proven incidents.'),
    ('05', 'Explain business impact', 'Translate technical findings into possible business consequences.',
     'Findings plus document context', 'Local Gemma groups affected areas and writes advisory explanations', 'Business impact context',
     'Gemma wording is clearly labeled advisory and does not replace source evidence or change severity.'),
    ('06', 'Verify the evidence', 'Keep each finding tied to evidence and assumptions.',
     'Potential finding and source references', 'Show exact location, confidence, attack path, and coverage limits', 'Evidence trail for review',
     'The interface exposes the source references and the limits of the supported analysis.'),
    ('07', 'Human review and reporting', 'Let a reviewer decide what should happen next.',
     'Evidence-backed candidate', 'Confirm, dismiss, or keep open with a review note', 'Recorded decision and exportable report',
     'This stage remains human-led. The product supports the decision; it does not make it automatically.'),
]

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=LINE, size='8'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)

def set_cell_margin(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)

def add_text(cell, text, bold=False, size=9.5, color=INK, align=None):
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return p

def add_para(doc, text='', style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.27
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        text = text[len(bold_lead):]
    if text:
        p.add_run(text)
    for r in p.runs:
        r.font.name = 'Arial'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor.from_string(INK)
    return p

def add_bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(.22)
        p.paragraph_format.first_line_indent = Inches(-.16)
        r = p.add_run(value)
        r.font.name = 'Arial'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        r.font.size = Pt(10.25)
        r.font.color.rgb = RGBColor.from_string(INK)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    r.font.color.rgb = RGBColor.from_string(INK)
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    r.italic = True
    r.font.name = 'Arial'
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MUTED)

def add_page_break(doc):
    doc.add_page_break()

def add_screenshot(doc, image_name, caption, width=6.75):
    path = ROOT / image_name
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)

def draw_rounded(draw, box, fill, outline=None, radius=22, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=pc(fill), outline=pc(outline) if outline else None, width=width)

def text_box(draw, box, title, body, number=None):
    x1, y1, x2, y2 = box
    draw_rounded(draw, box, WHITE, LINE, 20, 2)
    if number:
        draw.ellipse((x1+18, y1+18, x1+62, y1+62), fill=pc(INDIGO))
        draw.text((x1+40, y1+40), number, font=ImageFont.truetype(str(FONT_BOLD), 20), fill=pc(WHITE), anchor='mm')
    tx = x1 + (72 if number else 22)
    draw.text((tx, y1+20), title, font=ImageFont.truetype(str(FONT_BOLD), 25), fill=pc(INK))
    # Simple width-constrained manual wrapping.
    words = body.split()
    lines, line = [], ''
    for word in words:
        proposed = (line + ' ' + word).strip()
        if draw.textlength(proposed, font=ImageFont.truetype(str(FONT), 23)) > (x2-x1-42):
            lines.append(line); line = word
        else:
            line = proposed
    if line: lines.append(line)
    y = y1 + 72
    body_font = ImageFont.truetype(str(FONT), 23)
    for line in lines[:4]:
        draw.text((x1+22, y), line, font=body_font, fill=pc(MUTED))
        y += 31

def make_phase_diagram(number, title, subtitle, source, action, outcome):
    img = Image.new('RGB', (1600, 610), pc(WHITE))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1600, 610), fill='#F8F9FC')
    d.text((68, 45), f'{number}  {title}', font=ImageFont.truetype(str(FONT_BOLD), 42), fill=pc(INK))
    d.text((68, 105), subtitle, font=ImageFont.truetype(str(FONT), 25), fill=pc(MUTED))
    boxes = [(68, 180, 520, 460), (574, 180, 1026, 460), (1080, 180, 1532, 460)]
    text_box(d, boxes[0], 'What comes in', source, '1')
    text_box(d, boxes[1], 'What Invaria does', action, '2')
    text_box(d, boxes[2], 'What the reviewer sees', outcome, '3')
    for x in (537, 1043):
        d.line((x, 320, x+23, 320), fill=pc(INDIGO), width=5)
        d.polygon([(x+23, 308), (x+23, 332), (x+44, 320)], fill=pc(INDIGO))
    path = ROOT / f'phase-{number}.png'
    img.save(path)
    return path

def make_overview_diagram():
    img = Image.new('RGB', (1600, 840), pc(WHITE))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1600, 840), fill='#F8F9FC')
    d.text((68, 45), 'Invaria security review flow', font=ImageFont.truetype(str(FONT_BOLD), 38), fill=pc(INK))
    d.text((68, 100), 'The product turns a codebase into a human review decision, with evidence kept visible at every stage.', font=ImageFont.truetype(str(FONT), 20), fill=pc(MUTED))
    cols = 4
    w, h, x0, y0, gx, gy = 334, 190, 68, 185, 45, 55
    for i, (num, title, subtitle, *_rest) in enumerate(PHASES):
        row, col = divmod(i, cols)
        x = x0 + col*(w+gx); y = y0 + row*(h+gy)
        fill = INDIGO if i == 0 else WHITE
        line = INDIGO if i == 0 else LINE
        draw_rounded(d, (x, y, x+w, y+h), fill, line, 22, 2)
        fg = WHITE if i == 0 else INK
        mg = 'E7E6FF' if i == 0 else MUTED
        d.text((x+26, y+23), num, font=ImageFont.truetype(str(FONT_BOLD), 17), fill=pc(mg))
        d.text((x+26, y+55), title, font=ImageFont.truetype(str(FONT_BOLD), 22), fill=pc(fg))
        words, lines, linebuf = subtitle.split(), [], ''
        f = ImageFont.truetype(str(FONT), 15)
        for word in words:
            candidate = (linebuf+' '+word).strip()
            if d.textlength(candidate, font=f) > w-52:
                lines.append(linebuf); linebuf=word
            else: linebuf=candidate
        lines.append(linebuf)
        yy=y+95
        for line in lines[:3]:
            d.text((x+26, yy), line, font=f, fill=pc(mg)); yy += 22
        if i < 6:
            if col < 3:
                d.line((x+w+10, y+h/2, x+w+30, y+h/2), fill=pc(INDIGO), width=4)
                d.polygon([(x+w+30, y+h/2-9),(x+w+30,y+h/2+9),(x+w+45,y+h/2)], fill=pc(INDIGO))
            elif row == 0:
                d.line((x+w/2, y+h+12, x+w/2, y+h+36), fill=pc(INDIGO), width=4)
                d.polygon([(x+w/2-9,y+h+36),(x+w/2+9,y+h+36),(x+w/2,y+h+51)], fill=pc(INDIGO))
    path = ROOT / 'seven-stage-flow.png'
    img.save(path)
    return path

def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(INK)
    for name, size in [('Title', 30), ('Subtitle', 13), ('Heading 1', 19), ('Heading 2', 13)]:
        st = doc.styles[name]
        st.font.name = 'Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(INK)
        st.font.bold = True
    title_ppr = doc.styles['Title']._element.get_or_add_pPr()
    title_border = title_ppr.find(qn('w:pBdr'))
    if title_border is not None:
        title_ppr.remove(title_border)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Invaria review brief  |  Confidential draft for presentation  |  ')
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    for r in footer.runs:
        r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

def add_status_table(doc):
    rows = [
        ('Repository ingestion', 'Working for TaskForge', 'The TaskForge repository has imported successfully. Broader repository testing remains a next step.'),
        ('Code understanding', 'Working with coverage limits', 'Supported structures identify routes, handlers, inputs, and data access. TaskForge maps 19 endpoints, while source coverage still needs expansion.'),
        ('Security graph', 'Demo ready', 'The graph links a request to its connected code and data path for focused investigation.'),
        ('Potential findings', 'Demo ready', 'Static checks produce source-linked candidates that need human review.'),
        ('Gemma business impact', 'Connected locally', 'The local model produces advisory business context and affected-area labels.'),
        ('Evidence and review', 'Demo ready', 'Finding details retain references, assumptions, and review actions.'),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    headers = ['Area', 'Status', 'What this means for tomorrow']
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY); set_cell_border(cell, NAVY); set_cell_margin(cell)
        add_text(cell, value, True, 9.5, WHITE)
    set_repeat_table_header(table.rows[0])
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], 'F8F9FC' if idx % 2 else WHITE)
            set_cell_border(cells[i]); set_cell_margin(cells[i])
            add_text(cells[i], value, i == 0, 9.2)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    widths = [1.35, 1.35, 3.85]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
    return table

def build():
    diagrams = [make_phase_diagram(*p[:6]) for p in PHASES]
    overview = make_overview_diagram()
    doc = Document(); configure_document(doc)

    title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run('Invaria Security Review Brief')
    subtitle = doc.add_paragraph(style='Subtitle')
    subtitle.add_run('TaskForge demonstration  |  Product progress, business impact, and review flow')
    subtitle.paragraph_format.space_after = Pt(24)
    add_heading(doc, 'Purpose', 1)
    add_para(doc, 'This brief explains what the Invaria demo does today, how a repository moves through the analysis flow, and why the results matter to a business reviewer. The main conclusion is that the TaskForge demonstration is ready to show: it imports the sample repository, maps supported code paths, surfaces evidence-backed candidates, and explains their possible business impact with the local Gemma model. Human review remains the final decision point.')
    add_heading(doc, 'What to show in the review', 1)
    add_bullets(doc, [
        'Open the TaskForge Executive summary to start with the business view: 16 potential risks, the affected areas, and the review call to action.',
        'Choose a heatmap area, then use the risk journey to show how one endpoint could lead to a business consequence.',
        'Open the finding to show evidence, assumptions, source references, local Gemma context, and the confirm, dismiss, or keep-open decision.',
        'Switch to the security graph to show the route-to-data path behind the business summary.',
    ])
    add_heading(doc, 'Current demonstration status', 1)
    add_status_table(doc)

    add_page_break(doc)
    add_heading(doc, 'The current workspace', 1)
    add_para(doc, 'The Executive summary is the default completed-scan view. It turns the technical result into a plain-language discussion: what needs attention, which areas of the business could be affected, and which item the reviewer should inspect first.')
    add_screenshot(doc, '01-executive-dashboard.png', 'Figure 1. Current TaskForge Executive summary in the Invaria workspace.', 6.85)
    add_heading(doc, 'How to read this screen', 2)
    add_bullets(doc, [
        'The headline and action button identify the immediate review workload without claiming that a vulnerability is proven.',
        'Scorecards keep the scan size, business areas, high-priority routes, and review progress visible together.',
        'The heatmap and journey make the impact conversation accessible before a reviewer opens source-level detail.',
    ])

    add_page_break(doc)
    add_heading(doc, 'Why this matters in the real world', 1)
    add_para(doc, 'The TaskForge sample represents a decentralized compute marketplace. Its routes manage job creation, job results, lender profiles, marketplace listings, balances, and transaction-related actions. If an authorization or data-access weakness is confirmed in these paths, the consequence can extend beyond one technical issue.')
    impact_table = doc.add_table(rows=1, cols=3); impact_table.style='Table Grid'; impact_table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(['Potential business area', 'If a weakness is confirmed', 'What a reviewer needs to decide']):
        c=impact_table.rows[0].cells[i]; set_cell_shading(c,NAVY); set_cell_border(c,NAVY); set_cell_margin(c); add_text(c,head,True,9.5,WHITE)
    set_repeat_table_header(impact_table.rows[0])
    impact_rows = [
        ('Customer and partner trust', 'Unauthorized viewing of job or user information can reduce confidence in the marketplace.', 'Does the route enforce ownership or a role check before returning data?'),
        ('Marketplace operations', 'Changing job status, results, or acceptance flows without the right permission can disrupt delivery and dispute handling.', 'Which user or service is allowed to perform this action?'),
        ('Data confidentiality', 'Job details, lender information, and transaction records may expose sensitive commercial information.', 'Which fields are sensitive and who should be able to retrieve them?'),
        ('Financial workflows', 'Balance, lender, or transaction-related paths may affect reconciliation and user confidence.', 'What additional controls or audit checks are required before release?'),
    ]
    for ri, row in enumerate(impact_rows):
        cells=impact_table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], 'F8F9FC' if ri%2 else WHITE); set_cell_border(cells[i]); set_cell_margin(cells[i]); add_text(cells[i],value,i==0,9.25); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in impact_table.rows:
        for cell, width in zip(row.cells,[1.5,2.7,2.35]): cell.width=Inches(width)
    add_para(doc, 'These are possible consequences of findings under review. They are not evidence that an incident occurred. Invaria deliberately separates advisory business wording from the exact source evidence used to support a finding.')
    add_screenshot(doc, '02-findings-and-business-context.png', 'Figure 2. Findings view with the product overview and local Gemma business-impact context.', 6.85)

    add_page_break(doc)
    add_heading(doc, 'The seven stage security review flow', 1)
    add_para(doc, 'The workflow follows one simple principle: every business-facing message should lead back to code evidence and then to a human decision.')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(overview), width=Inches(6.85))
    add_caption(doc, 'Figure 3. The complete Invaria flow from repository intake to human review.')
    add_heading(doc, 'Stage by stage details', 2)
    add_para(doc, 'The following pages explain the input, product action, reviewer output, and current scope for every stage.')

    for page_start in range(0, len(PHASES), 2):
        add_page_break(doc)
        for idx in range(page_start, min(page_start+2, len(PHASES))):
            num, title, subtitle, source, action, outcome, current = PHASES[idx]
            add_heading(doc, f'{num}  {title}', 1)
            add_para(doc, subtitle)
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(diagrams[idx]), width=Inches(6.85))
            add_caption(doc, f'Figure {4+idx}. {title} flow.')
            add_para(doc, current, bold_lead=None)

    add_page_break(doc)
    add_heading(doc, 'Evidence and the security graph', 1)
    add_para(doc, 'The security graph is the bridge between the Executive summary and detailed evidence. A reviewer can select a route and see the connected inputs, functions, and data access steps. This helps the conversation move from “why it matters” to “show me why this is being reviewed.”')
    add_screenshot(doc, '03-security-graph.png', 'Figure 11. Security graph view showing a supported route-to-code path.', 6.85)
    add_heading(doc, 'What has been delivered so far', 1)
    add_bullets(doc, [
        'A working TaskForge repository analysis flow with the current supported static checks.',
        'A workspace with repository history, results navigation, finding details, evidence references, graph exploration, and report export actions.',
        'A local Gemma business-impact workflow that uses uploaded documents when useful and clearly labels generated wording as advisory.',
        'An Executive summary designed for a non-technical review: scorecards, heatmap, risk journey, balanced severity language, and an explicit human review action.',
        'A refined responsive interface with keyboard focus states and subtle interaction feedback for the demo flow.',
    ])

    add_page_break(doc)
    add_heading(doc, 'Future outcome and next steps', 1)
    add_para(doc, 'The immediate goal after the review is to broaden the evidence coverage while preserving the same review-first experience. The product should become more useful as it understands more languages, frameworks, document types, and application-specific authorization patterns.')
    next_table = doc.add_table(rows=1, cols=3); next_table.style='Table Grid'; next_table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(['Next outcome', 'Why it helps', 'Measure of progress']):
        c=next_table.rows[0].cells[i]; set_cell_shading(c,NAVY); set_cell_border(c,NAVY); set_cell_margin(c); add_text(c,head,True,9.5,WHITE)
    set_repeat_table_header(next_table.rows[0])
    next_rows = [
        ('Broader source coverage', 'More files and supported framework patterns make the analysis represent more of a real application.', 'Share of source files parsed and endpoints mapped.'),
        ('Richer business context', 'Better documents and curated business areas make the Gemma explanations more specific and useful.', 'Reviewer feedback on relevance and clarity.'),
        ('More precise checks', 'Additional evidence rules reduce noise and reveal more meaningful patterns.', 'Confirmed versus dismissed finding ratio.'),
        ('Review workflow maturity', 'Saved decisions and exports support collaboration, follow-up, and an audit trail.', 'Time from finding to documented decision.'),
    ]
    for ri,row in enumerate(next_rows):
        cells=next_table.add_row().cells
        for i,value in enumerate(row):
            set_cell_shading(cells[i], 'F8F9FC' if ri%2 else WHITE); set_cell_border(cells[i]); set_cell_margin(cells[i]); add_text(cells[i],value,i==0,9.25); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in next_table.rows:
        for cell,width in zip(row.cells,[1.6,3.0,1.95]): cell.width=Inches(width)
    add_heading(doc, 'Recommended closing message for the review', 1)
    add_para(doc, 'Invaria already provides a working path from a TaskForge repository to a business-facing review screen. The next investment is not a new dashboard; it is deeper source coverage and more precise evidence so that the same workflow can support a wider range of real applications with increasing confidence.')

    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()
