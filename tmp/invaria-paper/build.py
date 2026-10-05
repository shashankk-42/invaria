from pathlib import Path
from copy import deepcopy
import json,re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
TMP=ROOT/'tmp/invaria-paper'
data=json.loads((TMP/'results.json').read_text())
FONT='C:/Windows/Fonts/arial.ttf'
BOLD='C:/Windows/Fonts/arialbd.ttf'
def font(n,b=False):return ImageFont.truetype(BOLD if b else FONT,n)
def canvas(title,subtitle):
 im=Image.new('RGB',(1100,700),'white');dr=ImageDraw.Draw(im)
 dr.text((40,24),title,font=font(34,True),fill='#111111')
 dr.text((40,74),subtitle,font=font(25),fill='#444444')
 return im,dr
def txt(dr,xy,s,n=26,c='#111111',anchor=None,b=False):dr.text(xy,s,font=font(n,b),fill=c,anchor=anchor)
blue='#24557B';gray='#8396A5'
im,dr=canvas('Vulnerable and corrected fixtures','Observed candidates per fixture  |  6 independent cases')
left,right,top,bottom=110,1050,180,570
for v in [0,1]:
 y=bottom-v*300;dr.line((left,y,right,y),fill='#D6DCE1',width=2);txt(dr,(left-22,y),str(v),anchor='rm')
for j,label in enumerate(['Authorization','SQL injection','Credentials']):
 cx=270+j*320
 dr.rectangle((cx-92,270,cx-22,bottom),fill=blue)
 dr.line((cx+22,bottom,cx+92,bottom),fill=gray,width=8)
 txt(dr,(cx-57,242),'1',anchor='mm',b=True);txt(dr,(cx+57,542),'0',anchor='mm',b=True)
 txt(dr,(cx,bottom+43),label,anchor='mm')
dr.rectangle((250,635,280,660),fill=blue);txt(dr,(295,635),'Vulnerable',24)
dr.rectangle((585,635,615,660),fill=gray);txt(dr,(630,635),'Corrected',24)
im.save(TMP/'figure1.png',dpi=(300,300))
im,dr=canvas('Static engine analysis time','Median and p95  |  30 warmed calls per fixture')
left,right,top,bottom=340,1010,190,590
for v in [0,.2,.4,.6,.8]:
 x=left+(right-left)*v/.8;dr.line((x,top-25,x,bottom+15),fill='#D6DCE1',width=2);txt(dr,(x,bottom+42),f'{v:.1f}',24,anchor='mm')
for i,c in enumerate(data['cases']):
 y=top+i*75;name=c['case'].replace('bola','Authorization').replace('injection','Injection').replace('secrets','Credential').replace('-',' ')
 txt(dr,(left-20,y+10),name,24,anchor='rm')
 med=c['median_ms'];p95=c['p95_ms'];x=left+(right-left)*med/.8;xp=left+(right-left)*p95/.8
 dr.rectangle((left,y-6,x,y+26),fill=blue)
 dr.line((x,y+10,xp,y+10),fill='#111111',width=3);dr.line((xp,y,xp,y+20),fill='#111111',width=3)
 txt(dr,(x+12,y-27),f'{med:.3f}',22)
txt(dr,(675,673),'Elapsed time in milliseconds',25,anchor='mm')
im.save(TMP/'figure2.png',dpi=(300,300))
im,dr=canvas('Archived TaskForge scan coverage','8 September 2026  |  Candidates require human review')
left,right,top,bottom=110,1050,170,560
for v in [0,5,10,15,20]:
 y=bottom-v*18;dr.line((left,y,right,y),fill='#D6DCE1',width=2);txt(dr,(left-20,y),str(v),24,anchor='rm')
for j,(label,v,col) in enumerate([('Discovered\nroutes',19,blue),('Resolved\nroutes',19,blue),('Authorization\ncandidates',16,gray)]):
 cx=270+j*320;y=bottom-v*18;dr.rectangle((cx-70,y,cx+70,bottom),fill=col);txt(dr,(cx,y-25),str(v),30,anchor='mm',b=True)
 for k,line in enumerate(label.split('\n')):txt(dr,(cx,bottom+40+k*32),line,26,anchor='mm')
txt(dr,(550,677),'6 source files ingested  |  1 recorded as parsed',24,anchor='mm')
im.save(TMP/'figure3.png',dpi=(300,300))

# Start with the retained reference package and reuse its real style system.
d=Document(ROOT/'tmp/reference-paper.docx')
title_pattern=deepcopy(d.paragraphs[0]._p.pPr)
body_pattern=deepcopy(d.paragraphs[37]._p.pPr)
heading_pattern=deepcopy(d.paragraphs[36]._p.pPr)
first_sect=deepcopy(d.sections[0]._sectPr)
body_sect=deepcopy(d.sections[3]._sectPr)
for child in list(d._element.body):d._element.body.remove(child)
d._element.body.append(first_sect)
for ref in list(first_sect):
 if ref.tag in [qn('w:footerReference'),qn('w:headerReference'),qn('w:titlePg')]:first_sect.remove(ref)
for sty in d.styles:
 if sty.type==1:
  sty.font.color.rgb=RGBColor(0,0,0)
  if sty.name in ['Title','Normal','Body Text','Abstract','Author']:sty.font.name='Times New Roman'
if 'Title' not in d.styles:
 d.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH).base_style=d.styles['Author']
d.styles['Title'].font.size=Pt(24)
d.styles['Title'].font.color.rgb=RGBColor(0,0,0)
def paragraph(text='',kind='body'):
 p=d.add_paragraph()
 pattern=heading_pattern if kind=='heading' else body_pattern
 p._p.insert(0,deepcopy(pattern))
 p.style='Abstract' if kind=='heading' else 'Body Text'
 pf=p.paragraph_format
 pf.space_before=Pt(0);pf.space_after=Pt(6);pf.line_spacing=1.0
 pf.first_line_indent=Inches(.2) if kind=='body' else Inches(0)
 pf.left_indent=Inches(0);pf.right_indent=Inches(0)
 pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if kind=='body' else WD_ALIGN_PARAGRAPH.LEFT
 pf.keep_with_next=kind=='heading';pf.widow_control=True
 if kind=='heading':pf.space_before=Pt(8)
 r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(10);r.font.color.rgb=RGBColor(0,0,0);r.bold=kind=='heading'
 return p
lines=(TMP/'paper.md').read_text(encoding='utf-8').split('\n\n')
title=lines.pop(0).splitlines()[0].removeprefix('TITLE: ')
p=d.add_paragraph(title,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(14)
for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(24)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(6)
r=p.add_run('Shashank Kakad¹     Soham Moholkar     Omkar Kakade     Yash Pardeshi');r.font.name='Times New Roman';r.font.size=Pt(10)
# Affiliation is included only for the author whose details are supplied in the reference.
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10)
r=p.add_run('¹Vishwakarma Institute of Technology\nPune, India\nCorresponding author: shashank.kakad24@vit.edu');r.font.name='Times New Roman';r.font.size=Pt(9)
s=d.add_section(WD_SECTION_START.CONTINUOUS)
sp=s._sectPr
for child in list(sp):sp.remove(child)
for child in body_sect:sp.append(deepcopy(child))
for ref in list(sp):
 if ref.tag in [qn('w:footerReference'),qn('w:headerReference'),qn('w:titlePg')]:sp.remove(ref)
captions={1:'Fig. 1. Candidate counts in the three vulnerable/fixed pairs. Current evaluation, 3 October 2026.',2:'Fig. 2. Median engine time with p95 whiskers. In-memory analysis only; ingestion, storage and model review excluded.',3:'Fig. 3. Discovered routes, resolved routes and candidates from the archived TaskForge report. Counts represent different measures.'}
for block in lines:
 block=block.strip()
 if not block:continue
 if block.startswith('[FIGURE'):
  n=int(re.search(r'\d',block).group());p=paragraph('',kind='figure');p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(4);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.add_run().add_picture(str(TMP/f'figure{n}.png'),width=Inches(3.32))
  # Accessible image text retained in the embedded drawing.
  for x in p._p.xpath('.//wp:docPr'):x.set('descr',captions[n])
  p=paragraph(captions[n],kind='caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in p.runs:r.font.size=Pt(8);r.italic=True
 elif block.startswith('#'):
  paragraph(block.lstrip('#').strip(),'heading')
 else:
  p=paragraph(block)
  if block.startswith(('Abstract—','Keywords—')):
   p.paragraph_format.first_line_indent=Inches(0)
   lead,body=block.split('—',1);p.clear();r=p.add_run(lead+'—');r.bold=True;r.italic=True;p.add_run(body)
   for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
  if re.match(r'^\[\d+\]',block):
   p.paragraph_format.first_line_indent=Inches(-.18);p.paragraph_format.left_indent=Inches(.18)
   for r in p.runs:r.font.size=Pt(8)
   # Prevent long reference URLs from extending outside the column.
   for r in p.runs:r.text=re.sub(r'https://\S+',lambda m:m.group().replace('/','/\u200b').replace('-','-\u200b'),r.text)
d.core_properties.title=title;d.core_properties.author='Shashank Kakad; Soham Moholkar; Omkar Kakade; Yash Pardeshi'
d.core_properties.subject='Evidence based application security and invariant guided graph analysis'
d.core_properties.comments=''
out=ROOT/'research-output/INVARIA_Research_Paper.docx';d.save(out)
print('Saved',out,'words',sum(len(p.text.split()) for p in d.paragraphs),'sections',len(d.sections))
