from pathlib import Path
from shutil import copy2
from hashlib import sha256
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

source = Path(r"C:\Users\SHASHANK KAKAD\Downloads\FF_180__0 (2) copy.docx")
output = Path(r"C:\Users\SHASHANK KAKAD\Documents\projects\Invaria\artifacts\FF_180_Midsem_Review_1_Updated.docx")
output.parent.mkdir(parents=True, exist_ok=True)
print(f"SOURCE_SHA256={sha256(source.read_bytes()).hexdigest()}")
copy2(source, output)

doc = Document(output)
review_table = doc.tables[6]
header = review_table.cell(0, 0).paragraphs[0]
header.clear()
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = header.add_run("Review No.: 1                                  Group No.: TY-B5                             Date:")
run.bold = True
run.font.name = "Times New Roman"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
run.font.size = Pt(11)

body_cell = review_table.cell(1, 0)
for paragraph in body_cell.paragraphs:
    paragraph.clear()

title = body_cell.paragraphs[0]
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(7)
title_run = title.add_run("Progress Review Report")
title_run.bold = True
title_run.font.name = "Times New Roman"
title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
title_run.font.size = Pt(11)

entries = [
    ("Work completed to date:", True),
    ("1. Reviewed the OWASP Top 10, static application security testing approaches, and recent research on AI-assisted vulnerability detection. The study identified the limitation of code-only scanners in finding business-logic and database-schema-related vulnerabilities.", False),
    ("2. Finalized the problem statement, objectives, expected outcomes, and project milestones for the AI Agent for Code Vulnerability Detection.", False),
    ("3. Prepared the high-level architecture covering GitHub Pull Request ingestion, AST parsing using Tree-sitter, database-schema extraction, Neo4j knowledge graph construction, taint-path analysis, LLM-based reasoning, and reporting through GitHub review comments and a dashboard.", False),
    ("4. Selected the proposed technology stack: FastAPI for the backend, PostgreSQL and Redis for state and queues, Neo4j for the code-schema graph, Tree-sitter for code parsing, and Next.js for the findings dashboard.", False),
    ("5. Defined the initial implementation plan for repository webhook handling, source and migration-file parsing, and modelling API endpoints, functions, tables, and columns as graph entities.", False),
    ("Next phase: set up the development environment and begin the GitHub webhook and AST parsing prototype.", False),
]

for text, bold in entries:
    para = body_cell.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10.5)

# Remove unused blank slots in the source body area while preserving the row and table format.
for paragraph in body_cell.paragraphs[1:]:
    if not paragraph.text:
        p = paragraph._element
        p.getparent().remove(p)

# Keep document headers and footer content as in the reference; save only the copied document.
doc.save(output)
print(output)
