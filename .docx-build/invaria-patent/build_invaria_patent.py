from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph


REFERENCE = Path(r"C:\Users\SHASHANK KAKAD\Downloads\SoftGraph_Intelligence_Patent_Template_Filled_Updated.docx")
OUTPUT = Path(r"C:\Users\SHASHANK KAKAD\Documents\projects\Invaria\patent-output\Invaria_INVARIANT_Security_Reasoning_Patent_Draft.docx")
SYSTEM_ARCHITECTURE = Path(r"C:\Users\SHASHANK KAKAD\Documents\projects\Invaria\patent-output\assets\invaria-system-architecture-user.png")
ALGORITHM_WORKFLOW = Path(r"C:\Users\SHASHANK KAKAD\Documents\projects\Invaria\patent-output\assets\invariant-algorithm-workflow-user.png")


def replace_preserving_first_run(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def replace_all(paragraphs, old, new):
    for paragraph in paragraphs:
        if paragraph.text.strip() == old:
            replace_preserving_first_run(paragraph, new)
            return
    raise ValueError(f"Paragraph not found: {old}")


def remove_containing(paragraphs, needle):
    for paragraph in list(paragraphs):
        if needle.lower() in paragraph.text.lower():
            remove_paragraph(paragraph)


doc = Document(REFERENCE)
p = doc.paragraphs

# Keep the supplied applicant and inventor identity data, but replace the instruction-only text.
replace_preserving_first_run(p[0], "Patent Draft Information")
replace_all(p, "Copy and paste clear soft copy of signatures of all inventors in the following table: (Insert or delete the cells as per number of inventers)", "Signatures of inventors to be inserted before filing")
for needle in [
    "It must describe your invention in one sentence.",
    "Ex. which streams of science or technology",
    "What is the existing technology available",
    "Reason and advantages of your invention over existing technology",
    "Brief summary of your invention",
    "If any figures are given to describe the invention",
    "Describe in detail about the invention",
    "Preferably in one to two paragraphs",
    "Mention specific features of your invention",
    "Highlight what are the innovative features",
    "In which field or purpose your invention is useful",
    "Basic idea about your invention in one paragraph",
    "Drawings must not be labeled anywhere",
]:
    remove_containing(p, needle)

# The initial list is stale after removals; use exact old content to replace invention-specific slots.
p = doc.paragraphs
replacements = {
    "Hybrid Soft Computing System for Uncertainty-Aware Knowledge Graph Construction and Explainable Retrieval":
        "INVARIANT Based Security Reasoning for Software Repositories",
    "The invention relates to computer engineering, knowledge graphs, information retrieval, fuzzy logic, evolutionary computation, swarm intelligence, and explainable artificial intelligence.":
        "The invention relates to computer engineering, application security, static program analysis, data flow analysis, repository intelligence, and computer implemented security review systems.",
    "Existing knowledge-graph and retrieval systems commonly use hard confidence thresholds, static ranking weights, vector-only retrieval, or conventional graph traversal. These approaches may discard uncertain but useful relations, preserve noisy edges, require manual parameter tuning, or provide limited traceability. The proposed invention integrates fuzzy confidence, adaptive optimization, confidence-aware multi-hop graph search, semantic retrieval, and source provenance in one system. Specific patent and publication references shall be added after a formal prior-art search.":
        "Existing code assistants and conventional security scanners can generate or review code, identify known patterns, or report isolated findings. They may not construct an application level model of users, resources, actions, data flows, and authorization controls before assessing whether a security rule holds across a complete code path. The proposed system builds that model and applies an INVARIANT algorithm that derives and evaluates security rules against evidence from the repository. Specific patent and publication references shall be added after a formal prior art search.",
    "To construct confidence-weighted knowledge graphs from uncertain source material; adapt retrieval and graph-search parameters using evolutionary or swarm optimization; discover useful multi-hop graph paths using ant-colony search; and return results linked to source evidence, graph paths, and confidence information.":
        "To identify the security rules that must remain true in a software application, trace the code and data paths relevant to each rule, distinguish verified controls from possible gaps, and return review items with evidence that a developer can inspect.",
    "The system receives electronic source material, extracts entities and relations, assigns fuzzy confidence values, resolves equivalent entities, and stores confidence-weighted relations with provenance in a knowledge graph. For a query, semantic and graph retrieval are combined; ant-colony search explores multi-hop paths; evolutionary or swarm optimization tunes retrieval parameters; and the selected result is returned with supporting evidence and graph-path information.":
        "The system receives a software repository and scans source files, routes, APIs, data models, authorization checks, and external inputs. It constructs an application model and uses the INVARIANT algorithm to identify sensitive resources and actions, define security conditions for those resources and actions, trace relevant code paths, and record evidence for each condition. A decision engine returns a review item as verified, at risk, or requiring human review, together with the relevant repository locations and suggested remediation context.",
    "Figure 1 illustrates the overall architecture of the proposed system. Figure 2 illustrates knowledge-graph construction and adaptive feedback. Figure 3 illustrates query processing, graph-path exploration, ranking, and evidence generation.":
        "Figure 1 illustrates the Invaria system architecture. The system receives repository source, business context, and project policy inputs; performs secure ingestion and code extraction; constructs a security graph; executes the INVARIANT engine; verifies evidence; and records decisions. Security guidance, optional model review, and human feedback support the process, which produces scan records, a review workspace, and exported reports.\n\nFigure 2 illustrates the INVARIANT algorithm workflow. The workflow captures the repository, builds the application model, identifies a sensitive path, derives an invariant, traces relevant paths, collects and gates evidence, records a candidate when the required evidence is present, and routes the candidate for human review.",
    "110 denotes data-ingestion module; 120 preprocessing/chunking module; 130 entity and relation extraction module; 140 fuzzy uncertainty engine; 150 entity-resolution module; 160 knowledge-graph store; 161-165 graph nodes; 170 semantic vector store; 180 optimization engine; 190 ant-colony path-search module; 200 hybrid retrieval/ranking engine; 210 provenance/explainability engine; 220 query interface; 230 result output interface; 240 evaluation/feedback module; 290 parameter repository; 300 source-document repository; and 310 evidence package.":
        None,
    "Source material is ingested, segmented, and processed to extract candidate entities, relations, and supporting evidence. A fuzzy inference engine combines uncertainty signals such as extraction confidence, semantic similarity, source reliability, or graph consistency to assign a relation-confidence value. Entity resolution merges equivalent entities and the resulting nodes, typed relations, confidence values, and provenance are stored in a knowledge graph. During retrieval, semantic candidates and graph candidates are combined. An ant-colony module explores relevant multi-hop paths using edge-confidence and pheromone information, while an evolutionary or swarm optimizer tunes ranking weights, thresholds, path penalties, or related parameters. A hybrid ranker selects evidence and an explainability module returns the source location, graph path, and confidence information. No unverified numerical performance claim is made in this proposal.":
        "The repository scanner identifies entry points, request handlers, service functions, database operations, data models, authorization controls, configuration values, and externally controlled identifiers. The application model builder associates these elements with actors, resources, actions, and reachable paths. The model may use language-model or retrieval assistance to explain unfamiliar code, but the system records a finding only when repository evidence supports the corresponding decision.",
    "In a preferred embodiment, technical documents are split into source-linked units, entities and relations are extracted, fuzzy confidence values are assigned, duplicate entities are resolved, and a weighted knowledge graph is stored. On receiving a query, semantic retrieval identifies relevant evidence, ant-colony search explores graph paths, and an optimized hybrid ranker selects the strongest evidence. The output includes the source passage, selected graph path, and confidence information; evaluation feedback may subsequently be used to retune search parameters.":
        "In a preferred embodiment, a team connects a repository and identifies the application's relevant environment or framework. The system scans the repository, forms the application model, and selects sensitive resources and actions, such as customer records, transfers, refunds, tenant files, administrative actions, credentials, or payments. For each selected item, the INVARIANT algorithm creates a condition of the form: an actor may perform an action on a resource only when the required authentication, authorization, ownership, tenant, state, or policy condition holds.",
    "The inventive step is the coordinated use of fuzzy confidence during graph construction, confidence-aware ant-colony multi-hop search, evolutionary or swarm parameter optimization, and provenance-linked explainable retrieval. The same uncertainty information that influences graph-edge weighting also influences path exploration and final ranking, thereby reducing dependence on rigid thresholds and fixed manually chosen parameters while preserving traceability to source evidence.":
        "The inventive step is the INVARIANT algorithm. Rather than starting with a fixed vulnerability signature or a general code prompt, the algorithm derives application specific security rules from a repository model, identifies the actors, resources, actions, and policy conditions involved, and evaluates each rule along the relevant control and data paths. The same rule is connected to the evidence that supports a verified result, a potential violation, or an unresolved result. This produces a security review that remains tied to the application's own behavior and repository locations.",
    "The invention is useful in enterprise knowledge management, research and education platforms, technical-document search, compliance and policy navigation, scientific literature exploration, software documentation, cybersecurity knowledge systems, and other applications requiring retrieval from uncertain or heterogeneous information.":
        "The invention is useful for software teams developing web applications, software as a service products, financial technology products, healthcare systems, e-commerce platforms, enterprise applications, and AI assisted software projects that handle user data, money, credentials, or privileged actions.",
    "A computer-implemented system constructs and queries an uncertainty-aware knowledge graph from electronic source material. Extracted entities and relations are assigned fuzzy confidence values and stored with source provenance. Semantic retrieval is combined with graph retrieval, ant-colony search explores multi-hop paths, and evolutionary or swarm optimization adjusts retrieval parameters. A hybrid ranking engine selects relevant evidence, and an explainability engine returns the result together with supporting source evidence, graph-path information, and confidence values.":
        "A computer implemented security review system receives a software repository and constructs an application model containing repository entry points, code paths, data flows, resources, actions, and security controls. An INVARIANT algorithm derives a security condition for a selected resource or action, traces the corresponding code and data paths, and evaluates whether authentication, authorization, ownership, tenant, state, or policy conditions are evidenced along the path. The system records a verified result, a potential violation, or a result requiring human review, and outputs an evidence ledger containing repository locations, the relevant path, the affected condition, and remediation context.",
}

for old, new in replacements.items():
    if new is None:
        for paragraph in list(doc.paragraphs):
            if paragraph.text.strip() == old:
                remove_paragraph(paragraph)
                break
        else:
            raise ValueError(f"Paragraph not found: {old}")
    else:
        replace_all(doc.paragraphs, old, new)

# Add the core algorithmic detail as a second paragraph under detailed description.
p = doc.paragraphs
detail_anchor = next(x for x in p if x.text.startswith("The repository scanner identifies"))
detail_second = deepcopy(detail_anchor._element)
detail_anchor._element.addnext(detail_second)
detail_second_p = Paragraph(detail_second, detail_anchor._parent)
replace_preserving_first_run(
    detail_second_p,
    "For each candidate invariant, the algorithm creates a rule tuple comprising an actor, a resource, an action, and one or more required conditions. A representative rule requires that a requested action on a protected resource must be associated with an authenticated actor and a matching authorization, ownership, tenant, state, or policy check. The code and control flow analyzer traces the request from its entry point through service logic to a sensitive operation. The data flow analyzer determines whether a request controlled identifier or value reaches that operation. The evidence verifier compares the observed checks with the required conditions and records the supporting or missing evidence in the ledger.",
)

# Add a concrete preferred embodiment after the existing best method paragraph.
p = doc.paragraphs
best_anchor = next(x for x in p if x.text.startswith("In a preferred embodiment"))
best_second = deepcopy(best_anchor._element)
best_anchor._element.addnext(best_second)
best_second_p = Paragraph(best_second, best_anchor._parent)
replace_preserving_first_run(
    best_second_p,
    "For example, when an API receives an account identifier and retrieves a customer record, the algorithm traces the identifier from the route to the data access operation. It then seeks evidence that the requester is authenticated and that the customer record belongs to the requester or the requester’s tenant. If the path reaches the record without a matching check, the decision engine records the route, code locations, affected invariant, impact explanation, and a suggested review action.",
)

# Replace the template claims with concise Invaria claims and add one dependent claim.
claim1 = next(x for x in doc.paragraphs if x.text.startswith("1. A computer-implemented knowledge processing system"))
replace_preserving_first_run(
    claim1,
    "1. A computer implemented security review system comprising a repository scanner configured to obtain a software repository; an application model builder configured to identify entry points, code paths, data flows, resources, actions, and security controls; an INVARIANT engine configured to derive a security condition for a resource or action; a code and data flow analyzer configured to trace a path associated with the security condition; an evidence verifier; an evidence ledger; and a decision engine configured to produce a security review item associated with repository evidence.",
)
claim2 = next(x for x in doc.paragraphs if x.text.startswith("2. The system of claim 1"))
replace_preserving_first_run(
    claim2,
    "2. The system of claim 1, wherein the security condition comprises an actor, a resource, an action, and at least one required condition selected from authentication, authorization, ownership, tenant separation, state transition validity, credential protection, or policy compliance.",
)
claim3 = next(x for x in doc.paragraphs if x.text.startswith("3. The system of claim 1"))
replace_preserving_first_run(
    claim3,
    "3. The system of claim 1, wherein the code and data flow analyzer identifies whether an externally controlled identifier or value reaches a protected resource or sensitive operation without evidence of the required condition.",
)
claim4 = next(x for x in doc.paragraphs if x.text.startswith("4. A computer-implemented method"))
replace_preserving_first_run(
    claim4,
    "4. A computer implemented method comprising receiving a software repository; constructing an application model; identifying a protected resource or sensitive action; deriving an invariant comprising a required security condition; tracing a code and data path associated with the invariant; determining whether repository evidence supports the required security condition; recording a verified result, potential violation, or human review result; and outputting the result with repository locations and remediation context.",
)
claim4_element = claim4._element
new_claim = deepcopy(claim4_element)
claim4_element.addnext(new_claim)
new_claim_p = Paragraph(new_claim, claim4._parent)
replace_preserving_first_run(
    new_claim_p,
    "5. The method of claim 4, wherein a language model provides contextual assistance for interpreting source code, and wherein the evidence verifier requires code, control flow, data flow, configuration, or policy evidence before the decision engine records a result as verified.",
)

# Replace the generic source drawings with the two supplied Invaria figures.
for paragraph in list(doc.paragraphs):
    if (
        "graphic" in paragraph._p.xml
        or paragraph.text.strip() in {"Figure 1", "Figure 2", "Figure 3"}
    ):
        remove_paragraph(paragraph)

def add_figure(image_path, caption):
    figure = doc.add_paragraph()
    figure.alignment = 1
    figure.paragraph_format.keep_with_next = True
    figure.add_run().add_picture(str(image_path), width=Inches(6.5))
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.alignment = 1
    caption_paragraph.paragraph_format.keep_together = True
    caption_paragraph.runs[0].bold = True

add_figure(SYSTEM_ARCHITECTURE, "Figure 1")
add_figure(ALGORITHM_WORKFLOW, "Figure 2")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
