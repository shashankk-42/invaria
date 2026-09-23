from docx import Document
from invaria.documents import extract_text


def pdf_with_text(text: str) -> bytes:
    """Create a small valid PDF without adding a second PDF-writing dependency."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, object_body in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + object_body + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)


def test_extracts_docx_and_pdf_text(tmp_path):
    word = Document()
    word.add_heading("Architecture intent", 0)
    word.add_paragraph("Customers can inspect only their own compute jobs.")
    docx_path = tmp_path / "architecture.docx"
    word.save(docx_path)
    pdf_path = tmp_path / "api-intent.pdf"
    pdf_path.write_bytes(pdf_with_text("Only owners can read their job details"))

    assert "Customers can inspect" in extract_text(docx_path, ".docx")
    assert "Only owners can read" in extract_text(pdf_path, ".pdf")
